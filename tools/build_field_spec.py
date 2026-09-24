"""从字段清单 V2 导出领域包用的字段规格 JSON。

字段清单 Excel 是唯一的人类编辑入口，规则、渲染、摘要全部引用这里导出的 JSON。

用法（需使用带 openpyxl 的 Python）：
  python tools/build_field_spec.py

产物（可直接被 packages/field-spec 引用，禁止手工编辑）：
  packages/field-spec/src/field-spec.json        表单字段 191 项
  packages/field-spec/src/survey-checklist.json  量房确认清单 16 项
"""

import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_demo import BOOK, ROOT, build_spec  # noqa: E402

SPEC_OUT = ROOT / "packages" / "field-spec" / "src" / "field-spec.json"
SURVEY_OUT = ROOT / "packages" / "field-spec" / "src" / "survey-checklist.json"


def build_survey() -> list[dict]:
    """量房确认清单 16 项。

    机器可读映射（relatedFields / object / space / section / tier）是后补的资产，Excel 里没有这五列，
    所以重新生成时按序号从现有产物里保留——否则一跑构建就把这份映射冲掉，清单的合并与分区会静默退化。
    """
    ws = openpyxl.load_workbook(BOOK)["量房确认清单"]
    existing: dict[int, dict] = {}
    if SURVEY_OUT.exists():
        existing = {int(i["no"]): i for i in json.loads(SURVEY_OUT.read_text(encoding="utf-8"))}
    items = []
    for r in range(5, 40):
        no = ws.cell(r, 1).value
        item = ws.cell(r, 2).value
        if not isinstance(no, int) or not item:
            continue
        row = {
            "no": no,
            "item": item.strip(),
            "goal": (ws.cell(r, 3).value or "").strip(),
            "source": (ws.cell(r, 4).value or "").strip(),
        }
        previous = existing.get(no, {})
        for key in ("relatedFields", "object", "space", "section", "tier"):
            if key in previous:
                row[key] = previous[key]
        items.append(row)
    return items


def dump(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    print(f"生成 {path.relative_to(ROOT)}：{len(payload)} 项")


def main() -> None:
    spec = build_spec()
    survey = build_survey()
    sections: dict[str, int] = {}
    for f in spec:
        sections[f["section"]] = sections.get(f["section"], 0) + 1
    dump(SPEC_OUT, spec)
    dump(SURVEY_OUT, survey)
    print(f"分区 {len(sections)} 个：{sections}")


if __name__ == "__main__":
    main()
