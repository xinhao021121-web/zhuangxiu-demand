"""从字段清单 V2 生成 Demo 页面：把字段规格注入 tools/_demo_template.html。

用法（需使用内置 Python，因其带有 openpyxl）：
  python tools/build_demo.py
模板：tools/_demo_template.html（勿直接打开，缺少字段数据时只会显示提示）
产物：demo/装修需求发现助手_Demo_V0.2.html
"""

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = (
    ROOT
    / "outputs"
    / "01a0c95f-3b4f-7762-ad52-61bb13941fcc"
    / "装修需求采集表_字段清单_V2.xlsx"
)
TEMPLATE = ROOT / "tools" / "_demo_template.html"
OUTPUT = ROOT / "demo" / "装修需求发现助手_Demo_V0.2.html"
PLACEHOLDER = "/*__FIELD_SPEC__*/[]"

SELECT_TYPES = {"单选", "多选"}


def build_spec() -> list[dict]:
    ws = openpyxl.load_workbook(BOOK)["表单字段"]
    spec = []
    for r in range(5, 300):
        fid = ws.cell(r, 1).value
        if not isinstance(fid, str) or not re.match(r"^[a-z][a-z0-9_]*$", fid):
            continue
        ftype = (ws.cell(r, 5).value or "文本").strip()
        raw = (ws.cell(r, 6).value or "").strip()
        options = (
            [o.strip() for o in re.split(r"\s*/\s*", raw) if o.strip()]
            if ftype in SELECT_TYPES
            else []
        )
        spec.append(
            {
                "id": fid,
                "section": ws.cell(r, 2).value,
                "group": ws.cell(r, 3).value,
                "label": ws.cell(r, 4).value,
                "type": ftype,
                "unit": "" if ftype in SELECT_TYPES else raw,
                "options": options,
                "suggest": (ws.cell(r, 7).value or "选填").strip(),
                "scope": (ws.cell(r, 8).value or "固定").strip(),
                "applies": (ws.cell(r, 9).value or "全部").strip(),
            }
        )
    return spec


def main() -> None:
    spec = build_spec()
    payload = json.dumps(spec, ensure_ascii=False, separators=(",", ":"))
    html = TEMPLATE.read_text(encoding="utf-8")
    if PLACEHOLDER not in html:
        raise SystemExit("模板中未找到字段占位符")
    OUTPUT.write_text(html.replace(PLACEHOLDER, payload), encoding="utf-8")
    scopes = {}
    for f in spec:
        scopes[f["scope"]] = scopes.get(f["scope"], 0) + 1
    print(f"生成 {OUTPUT.name}：{len(spec)} 个字段，{OUTPUT.stat().st_size} 字节，空间类型 {scopes}")


if __name__ == "__main__":
    main()
